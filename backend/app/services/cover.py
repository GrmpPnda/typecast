from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from app import paths

UPLOAD_DIR = paths.UPLOAD_DIR

DPI = 300
BLEED = 0.125  # inches


@dataclass
class CoverDimensions:
    trim_width: float
    trim_height: float
    page_count: int
    spine_width: float
    bleed: float
    total_width: float
    total_height: float
    front_x: float
    spine_x: float
    back_x: float
    cover_type: str = "paperback"
    flap_width: float = 0.0
    front_flap_x: float = 0.0
    back_flap_x: float = 0.0

    def to_dict(self) -> dict:
        d = {
            "trim_width": self.trim_width,
            "trim_height": self.trim_height,
            "page_count": self.page_count,
            "spine_width": round(self.spine_width, 4),
            "bleed": self.bleed,
            "total_width": round(self.total_width, 4),
            "total_height": round(self.total_height, 4),
            "front_x": round(self.front_x, 4),
            "spine_x": round(self.spine_x, 4),
            "back_x": round(self.back_x, 4),
            "total_width_px": round(self.total_width * DPI),
            "total_height_px": round(self.total_height * DPI),
            "cover_type": self.cover_type,
        }
        if self.cover_type == "hardcover":
            d["flap_width"] = round(self.flap_width, 4)
            d["front_flap_x"] = round(self.front_flap_x, 4)
            d["back_flap_x"] = round(self.back_flap_x, 4)
        return d


def calculate_dimensions(
    trim_width: float,
    trim_height: float,
    page_count: int,
    spine_factor: float = 0.0025,
    cover_type: str = "paperback",
    flap_width: float = 3.5,
) -> CoverDimensions:
    spine_width = page_count * spine_factor

    if cover_type == "hardcover":
        # Dust jacket: back_flap + back + spine + front + front_flap (all with bleed)
        panel_w = trim_width + BLEED
        total_width = (flap_width + BLEED) + panel_w + spine_width + panel_w + (flap_width + BLEED)
        total_height = trim_height + 2 * BLEED

        back_flap_x = 0.0
        back_x = flap_width + BLEED
        spine_x = back_x + trim_width + BLEED
        front_x = spine_x + spine_width
        front_flap_x = front_x + trim_width + BLEED

        return CoverDimensions(
            trim_width=trim_width,
            trim_height=trim_height,
            page_count=page_count,
            spine_width=spine_width,
            bleed=BLEED,
            total_width=total_width,
            total_height=total_height,
            front_x=front_x,
            spine_x=spine_x,
            back_x=back_x,
            cover_type=cover_type,
            flap_width=flap_width,
            front_flap_x=front_flap_x,
            back_flap_x=back_flap_x,
        )

    # Paperback
    total_width = (trim_width + BLEED) + spine_width + (trim_width + BLEED)
    total_height = trim_height + 2 * BLEED

    back_x = 0.0
    spine_x = trim_width + BLEED
    front_x = spine_x + spine_width

    return CoverDimensions(
        trim_width=trim_width,
        trim_height=trim_height,
        page_count=page_count,
        spine_width=spine_width,
        bleed=BLEED,
        total_width=total_width,
        total_height=total_height,
        front_x=front_x,
        spine_x=spine_x,
        back_x=back_x,
        cover_type=cover_type,
    )


def _inches_to_px(inches: float) -> int:
    return round(inches * DPI)


def _load_and_fit(image_path: str, target_w: int, target_h: int) -> Image.Image:
    img = Image.open(image_path).convert("RGB")
    img_ratio = img.width / img.height
    target_ratio = target_w / target_h
    if img_ratio > target_ratio:
        new_h = target_h
        new_w = round(target_h * img_ratio)
    else:
        new_w = target_w
        new_h = round(target_w / img_ratio)
    img = img.resize((new_w, new_h), Image.LANCZOS)
    left = (new_w - target_w) // 2
    top = (new_h - target_h) // 2
    return img.crop((left, top, left + target_w, top + target_h))


def _parse_css_font_family(css_value: str) -> list[str]:
    names: list[str] = []
    for part in css_value.split(","):
        name = part.strip().strip("'\"")
        if name and name.lower() not in ("serif", "sans-serif", "monospace", "sans serif"):
            names.append(name)
    return names


def _resolve_font_path(family: str | None) -> str | None:
    if not family:
        return None
    candidates = _parse_css_font_family(family)
    fonts_dir = UPLOAD_DIR / "fonts"
    if fonts_dir.exists():
        for name in candidates:
            nl = name.lower()
            for f in fonts_dir.iterdir():
                if f.suffix.lower() in (".ttf", ".otf", ".woff", ".woff2"):
                    try:
                        loaded = ImageFont.truetype(str(f), 20)
                        family, style = loaded.getname()
                        family_l = family.lower()
                        full_l = f"{family} {style}".lower().strip()
                        if nl == family_l or nl == full_l or nl in family_l or family_l in nl:
                            return str(f)
                    except Exception:
                        continue

    for name in candidates:
        try:
            ImageFont.truetype(name, 20)
            return name
        except OSError:
            continue
    return None


def _load_font(path: str | None, size_pt: int) -> ImageFont.FreeTypeFont:
    if path:
        try:
            return ImageFont.truetype(path, size_pt)
        except OSError:
            pass
    try:
        return ImageFont.truetype("Arial", size_pt)
    except OSError:
        return ImageFont.load_default()


def _find_font(family: str | None, size_pt: int) -> ImageFont.FreeTypeFont:
    return _load_font(_resolve_font_path(family), size_pt)


def _draw_centered_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    font: ImageFont.FreeTypeFont,
    area: tuple[int, int, int, int],
    fill: str = "#FFFFFF",
):
    x1, y1, x2, y2 = area
    bbox = draw.textbbox((0, 0), text, font=font)
    tw = bbox[2] - bbox[0]
    th = bbox[3] - bbox[1]
    tx = x1 + (x2 - x1 - tw) // 2
    ty = y1 + (y2 - y1 - th) // 2
    draw.text((tx, ty), text, font=font, fill=fill)


def _draw_vertical_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    font: ImageFont.FreeTypeFont,
    area: tuple[int, int, int, int],
    fill: str = "#FFFFFF",
):
    x1, y1, x2, y2 = area
    temp = Image.new("RGBA", (y2 - y1, x2 - x1), (0, 0, 0, 0))
    temp_draw = ImageDraw.Draw(temp)
    bbox = temp_draw.textbbox((0, 0), text, font=font)
    tw = bbox[2] - bbox[0]
    th = bbox[3] - bbox[1]
    tx = (temp.width - tw) // 2
    ty = (temp.height - th) // 2
    temp_draw.text((tx, ty), text, font=font, fill=fill)
    rotated = temp.rotate(90, expand=True)
    paste_x = x1 + (x2 - x1 - rotated.width) // 2
    paste_y = y1 + (y2 - y1 - rotated.height) // 2
    return rotated, (paste_x, paste_y)


def _wrap_text(
    text: str, font: ImageFont.FreeTypeFont, max_width: int, draw: ImageDraw.ImageDraw
) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        test = f"{current} {word}".strip()
        bbox = draw.textbbox((0, 0), test, font=font)
        if bbox[2] - bbox[0] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def generate_full_cover(
    dims: CoverDimensions,
    front_image_path: str | None = None,
    back_image_path: str | None = None,
    spine_text: str = "",
    back_title: str = "",
    back_blurb: str = "",
    background_color: str = "#1a1a2e",
    text_color: str = "#FFFFFF",
    spine_font_family: str | None = None,
    blurb_font_family: str | None = None,
    barcode_zone: bool = True,
    back_overlay_opacity: int = 180,
    draw_guides: bool = False,
    back_logo_path: str | None = None,
    back_website: str = "",
    front_flap_text: str = "",
    back_flap_text: str = "",
) -> BytesIO:
    w_px = _inches_to_px(dims.total_width)
    h_px = _inches_to_px(dims.total_height)

    canvas = Image.new("RGB", (w_px, h_px), background_color)

    front_w = _inches_to_px(dims.trim_width + BLEED)
    front_h = h_px
    spine_w = _inches_to_px(dims.spine_width)
    back_w = front_w

    front_x = _inches_to_px(dims.front_x)
    spine_x = _inches_to_px(dims.spine_x)

    back_x = _inches_to_px(dims.back_x)

    if back_image_path:
        resolved = _resolve_image_path(back_image_path)
        if resolved:
            back_img = _load_and_fit(resolved, back_w, h_px)
            canvas.paste(back_img, (back_x, 0))

    if front_image_path:
        resolved = _resolve_image_path(front_image_path)
        if resolved:
            front_img = _load_and_fit(resolved, front_w, front_h)
            canvas.paste(front_img, (front_x, 0))

    draw = ImageDraw.Draw(canvas)

    if spine_text and spine_w > _inches_to_px(0.08):
        spine_font_size = min(max(_inches_to_px(0.12), 20), spine_w - 10)
        spine_font = _find_font(spine_font_family, spine_font_size)
        rotated, pos = _draw_vertical_text(
            draw, spine_text, spine_font,
            (spine_x, _inches_to_px(BLEED), spine_x + spine_w, h_px - _inches_to_px(BLEED)),
            fill=text_color,
        )
        canvas.paste(rotated, pos, rotated)

    _draw_back_cover_text(
        canvas, draw, dims, back_w,
        back_title=back_title,
        back_blurb=back_blurb,
        text_color=text_color,
        blurb_font_family=blurb_font_family,
        barcode_zone=barcode_zone,
        has_back_image=back_image_path is not None,
        overlay_opacity=back_overlay_opacity,
        back_logo_path=back_logo_path,
        back_website=back_website,
        back_x_offset=back_x,
    )

    if dims.cover_type == "hardcover":
        _draw_flap_text(
            canvas, draw, dims,
            front_flap_text=front_flap_text,
            back_flap_text=back_flap_text,
            text_color=text_color,
            font_family=blurb_font_family,
        )

    if draw_guides:
        _draw_guides(draw, dims, w_px, h_px)

    buf = BytesIO()
    canvas.save(buf, format="PNG", dpi=(DPI, DPI))
    buf.seek(0)
    return buf


def _measure_blurb(
    paragraphs: list[str],
    font: ImageFont.FreeTypeFont,
    max_width: int,
    line_h: int,
    para_gap: int,
    draw: ImageDraw.ImageDraw,
) -> int:
    total = 0
    for i, para in enumerate(paragraphs):
        lines = _wrap_text(para, font, max_width, draw)
        total += len(lines) * line_h
        if i < len(paragraphs) - 1:
            total += para_gap
    return total


def _draw_back_cover_text(
    canvas: Image.Image,
    draw: ImageDraw.ImageDraw,
    dims: CoverDimensions,
    back_w: int,
    back_title: str,
    back_blurb: str,
    text_color: str,
    blurb_font_family: str | None,
    barcode_zone: bool,
    has_back_image: bool,
    overlay_opacity: int,
    back_logo_path: str | None = None,
    back_website: str = "",
    back_x_offset: int = 0,
):
    h_px = canvas.height
    bleed_px = _inches_to_px(BLEED)
    margin = _inches_to_px(0.5)
    ox = back_x_offset

    area_x1 = ox + bleed_px + margin
    area_x2 = ox + back_w - margin
    area_y1 = bleed_px + margin
    area_y2 = h_px - bleed_px - margin
    text_width = area_x2 - area_x1

    if not back_title and not back_blurb and not barcode_zone:
        return

    if has_back_image and (back_title or back_blurb):
        overlay = Image.new("RGBA", (back_w, h_px), (0, 0, 0, 0))
        overlay_draw = ImageDraw.Draw(overlay)
        overlay_draw.rectangle(
            [bleed_px, bleed_px, back_w - bleed_px, h_px - bleed_px],
            fill=(0, 0, 0, overlay_opacity),
        )
        canvas.paste(
            Image.alpha_composite(
                canvas.crop((ox, 0, ox + back_w, h_px)).convert("RGBA"),
                overlay,
            ).convert("RGB"),
            (ox, 0),
        )
        draw = ImageDraw.Draw(canvas)

    font_path = _resolve_font_path(blurb_font_family)

    barcode_bottom_y = area_y2
    if barcode_zone:
        barcode_w = _inches_to_px(1.5)
        barcode_h = _inches_to_px(1.125)
        barcode_margin = _inches_to_px(0.15)
        bx2 = ox + back_w - _inches_to_px(0.25)
        by2 = h_px - bleed_px - _inches_to_px(0.25)
        bx = bx2 - barcode_w
        by = by2 - barcode_h
        draw.rectangle([bx, by, bx2, by2], fill="#FFFFFF")
        small_font = _load_font(None, _inches_to_px(0.06))
        draw.text(
            (bx + barcode_margin, by + barcode_margin),
            "ISBN\nBarcode\nArea",
            font=small_font,
            fill="#999999",
        )
        barcode_bottom_y = by - 12

    y = area_y1

    title_font_size = _inches_to_px(0.18)
    title_font = _load_font(font_path, title_font_size)

    title_h = 0
    if back_title:
        title_lines = _wrap_text(back_title.upper(), title_font, text_width, draw)
        title_line_h = _inches_to_px(0.24)
        title_h = len(title_lines) * title_line_h + _inches_to_px(0.2)

    blurb_space = barcode_bottom_y - area_y1 - title_h
    paragraphs = [p.strip() for p in back_blurb.split("\n") if p.strip()] if back_blurb else []

    if paragraphs and blurb_space > 0:
        best_size = _inches_to_px(0.07)
        for trial_pt in range(_inches_to_px(0.16), _inches_to_px(0.06), -1):
            trial_font = _load_font(font_path, trial_pt)
            line_h = int(trial_pt * 1.5)
            para_gap = int(trial_pt * 0.8)
            needed = _measure_blurb(paragraphs, trial_font, text_width, line_h, para_gap, draw)
            if needed <= blurb_space:
                best_size = trial_pt
                break
        blurb_font = _load_font(font_path, best_size)
        blurb_line_h = int(best_size * 1.5)
        blurb_para_gap = int(best_size * 0.8)
    else:
        blurb_font = _load_font(font_path, _inches_to_px(0.1))
        blurb_line_h = _inches_to_px(0.15)
        blurb_para_gap = _inches_to_px(0.1)

    if back_title:
        title_line_h = _inches_to_px(0.24)
        for line in _wrap_text(back_title.upper(), title_font, text_width, draw):
            bbox = draw.textbbox((0, 0), line, font=title_font)
            tw = bbox[2] - bbox[0]
            tx = area_x1 + (text_width - tw) // 2
            draw.text((tx, y), line, font=title_font, fill=text_color)
            y += title_line_h
        y += _inches_to_px(0.2)

    for i, para in enumerate(paragraphs):
        if y >= barcode_bottom_y:
            break
        lines = _wrap_text(para, blurb_font, text_width, draw)
        for line in lines:
            if y + blurb_line_h > barcode_bottom_y:
                break
            draw.text((area_x1, y), line, font=blurb_font, fill=text_color)
            y += blurb_line_h
        if i < len(paragraphs) - 1:
            y += blurb_para_gap

    bottom_y = area_y2
    if back_logo_path:
        resolved = _resolve_image_path(back_logo_path)
        if resolved:
            logo_max_h = _inches_to_px(0.6)
            logo_max_w = _inches_to_px(1.2)
            try:
                logo = Image.open(resolved).convert("RGBA")
                ratio = min(logo_max_w / logo.width, logo_max_h / logo.height)
                new_w = round(logo.width * ratio)
                new_h = round(logo.height * ratio)
                logo = logo.resize((new_w, new_h), Image.LANCZOS)
                logo_x = area_x1
                logo_y = bottom_y - new_h
                canvas.paste(logo, (logo_x, logo_y), logo)
                bottom_y = logo_y - _inches_to_px(0.08)
            except Exception:
                pass

    if back_website:
        website_font = _load_font(font_path, _inches_to_px(0.08))
        bbox = draw.textbbox((0, 0), back_website, font=website_font)
        wy = bottom_y - (bbox[3] - bbox[1])
        draw.text((area_x1, wy), back_website, font=website_font, fill=text_color)


def _draw_flap_text(
    canvas: Image.Image,
    draw: ImageDraw.ImageDraw,
    dims: CoverDimensions,
    front_flap_text: str,
    back_flap_text: str,
    text_color: str,
    font_family: str | None,
):
    """Draw text on front and back flaps of a dust jacket."""
    if not front_flap_text and not back_flap_text:
        return

    bleed_px = _inches_to_px(BLEED)
    flap_w_px = _inches_to_px(dims.flap_width)
    margin = _inches_to_px(0.4)
    font_path = _resolve_font_path(font_family)
    flap_font = _load_font(font_path, _inches_to_px(0.1))
    line_h = _inches_to_px(0.15)
    text_w = flap_w_px - 2 * margin

    if front_flap_text:
        fx = _inches_to_px(dims.front_flap_x) + margin
        fy = bleed_px + margin
        for para in front_flap_text.split("\n"):
            if not para.strip():
                fy += line_h
                continue
            for line in _wrap_text(para.strip(), flap_font, text_w, draw):
                draw.text((fx, fy), line, font=flap_font, fill=text_color)
                fy += line_h
            fy += int(line_h * 0.5)

    if back_flap_text:
        bx = _inches_to_px(dims.back_flap_x) + margin
        by = bleed_px + margin
        for para in back_flap_text.split("\n"):
            if not para.strip():
                by += line_h
                continue
            for line in _wrap_text(para.strip(), flap_font, text_w, draw):
                draw.text((bx, by), line, font=flap_font, fill=text_color)
                by += line_h
            by += int(line_h * 0.5)


def _resolve_image_path(path: str) -> str | None:
    if path.startswith("/uploads/"):
        full = UPLOAD_DIR / path[len("/uploads/"):]
        return str(full) if full.exists() else None
    p = Path(path)
    if p.exists():
        return str(p)
    full = UPLOAD_DIR / path.lstrip("/")
    return str(full) if full.exists() else None


def _draw_guides(
    draw: ImageDraw.ImageDraw,
    dims: CoverDimensions,
    w_px: int,
    h_px: int,
):
    bleed_px = _inches_to_px(BLEED)
    guide_color = (255, 0, 0, 128)

    trim_lines = [
        (bleed_px, 0, bleed_px, h_px),
        (w_px - bleed_px, 0, w_px - bleed_px, h_px),
        (0, bleed_px, w_px, bleed_px),
        (0, h_px - bleed_px, w_px, h_px - bleed_px),
    ]

    spine_x = _inches_to_px(dims.spine_x)
    spine_w = _inches_to_px(dims.spine_width)
    trim_lines.extend([
        (spine_x, 0, spine_x, h_px),
        (spine_x + spine_w, 0, spine_x + spine_w, h_px),
    ])

    if dims.cover_type == "hardcover":
        back_x = _inches_to_px(dims.back_x)
        front_flap_x = _inches_to_px(dims.front_flap_x)
        trim_lines.extend([
            (back_x, 0, back_x, h_px),
            (front_flap_x, 0, front_flap_x, h_px),
        ])

    for line in trim_lines:
        draw.line(line, fill=guide_color, width=1)


def generate_cover_pdf(png_buf: BytesIO, dims: CoverDimensions) -> BytesIO:
    """Wrap the rendered cover PNG into a PDF at exact physical dimensions."""
    from reportlab.lib.units import inch
    from reportlab.pdfgen import canvas as pdf_canvas

    width_pts = dims.total_width * inch
    height_pts = dims.total_height * inch

    buf = BytesIO()
    c = pdf_canvas.Canvas(buf, pagesize=(width_pts, height_pts))
    png_buf.seek(0)

    from reportlab.lib.utils import ImageReader

    img = ImageReader(png_buf)
    c.drawImage(img, 0, 0, width=width_pts, height=height_pts)
    c.save()
    buf.seek(0)
    return buf


def generate_cover_docx(
    dims: CoverDimensions,
    front_image_path: str | None = None,
    back_image_path: str | None = None,
    spine_text: str = "",
    back_title: str = "",
    back_blurb: str = "",
    background_color: str = "#1a1a2e",
    text_color: str = "#FFFFFF",
    spine_font_family: str | None = None,
    blurb_font_family: str | None = None,
    back_website: str = "",
    front_flap_text: str = "",
    back_flap_text: str = "",
) -> BytesIO:
    """Generate an editable DOCX of the cover layout using a wide landscape page."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(dims.total_width)
    section.page_height = Inches(dims.total_height)
    section.top_margin = Inches(BLEED)
    section.bottom_margin = Inches(BLEED)
    section.left_margin = Inches(BLEED)
    section.right_margin = Inches(BLEED)

    text_rgb = _hex_to_rgb(text_color)

    if dims.cover_type == "hardcover":
        col_widths = [
            dims.flap_width,
            dims.trim_width,
            dims.spine_width,
            dims.trim_width,
            dims.flap_width,
        ]
        headers = ["Back Flap", "Back Cover", "Spine", "Front Cover", "Front Flap"]
    else:
        col_widths = [dims.trim_width, dims.spine_width, dims.trim_width]
        headers = ["Back Cover", "Spine", "Front Cover"]

    table = doc.add_table(rows=2, cols=len(col_widths))
    table.autofit = False

    for i, w in enumerate(col_widths):
        for row in table.rows:
            row.cells[i].width = Inches(w)

    header_row = table.rows[0]
    for i, hdr in enumerate(headers):
        cell = header_row.cells[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(hdr)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(150, 150, 150)

    content_row = table.rows[1]
    content_row.height = Inches(dims.trim_height - 0.5)

    if dims.cover_type == "hardcover":
        _docx_cell_text(
            content_row.cells[0], back_flap_text, text_rgb, blurb_font_family,
        )
        _docx_back_cover(
            content_row.cells[1], back_title, back_blurb, back_website,
            text_rgb, blurb_font_family, back_image_path,
        )
        _docx_spine(content_row.cells[2], spine_text, text_rgb, spine_font_family)
        _docx_front_cover(content_row.cells[3], front_image_path)
        _docx_cell_text(
            content_row.cells[4], front_flap_text, text_rgb, blurb_font_family,
        )
    else:
        _docx_back_cover(
            content_row.cells[0], back_title, back_blurb, back_website,
            text_rgb, blurb_font_family, back_image_path,
        )
        _docx_spine(content_row.cells[1], spine_text, text_rgb, spine_font_family)
        _docx_front_cover(content_row.cells[2], front_image_path)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _hex_to_rgb(hex_color: str) -> tuple[int, int, int]:
    h = hex_color.lstrip("#")
    return (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _docx_cell_text(
    cell, text: str, rgb: tuple[int, int, int], font_family: str | None,
):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    cell.paragraphs[0].clear()
    for para_text in (text or "").split("\n"):
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(para_text)
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(*rgb)
        if font_family:
            run.font.name = _parse_css_font_family(font_family)[0]


def _docx_back_cover(
    cell, title: str, blurb: str, website: str,
    rgb: tuple[int, int, int], font_family: str | None,
    back_image_path: str | None,
):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    cell.paragraphs[0].clear()
    font_name = _parse_css_font_family(font_family)[0] if font_family else None

    if back_image_path:
        resolved = _resolve_image_path(back_image_path)
        if resolved:
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(resolved, height=Inches(2.0))

    if title:
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title.upper())
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(*rgb)
        if font_name:
            run.font.name = font_name

    if blurb:
        for para_text in blurb.split("\n"):
            if not para_text.strip():
                cell.add_paragraph()
                continue
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(para_text.strip())
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(*rgb)
            if font_name:
                run.font.name = font_name

    if website:
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(website)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(*rgb)


def _docx_spine(
    cell, text: str, rgb: tuple[int, int, int], font_family: str | None,
):
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.paragraphs[0].clear()
    if text:
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(*rgb)
        if font_family:
            run.font.name = _parse_css_font_family(font_family)[0]


def _docx_front_cover(cell, front_image_path: str | None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    cell.paragraphs[0].clear()
    if front_image_path:
        resolved = _resolve_image_path(front_image_path)
        if resolved:
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(resolved, height=Inches(4.0))
